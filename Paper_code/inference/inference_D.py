"""
Inference — Config D : NLLB-200 + LoRA + Back-Translation.
Traduit 20 exemples du split test → outputs/resultats_config_D.docx
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

os.environ["KMP_DUPLICATE_LIB_OK"] = "TRUE"
os.environ["TRANSFORMERS_SAFE_LOAD"] = "1"

import json
import torch
from docx import Document
from config import BackTranslationConfig
from reproducibility import best_checkpoint
from models.model_lora import NLLBFineTuner

cfg = BackTranslationConfig()
os.makedirs(cfg.OUTPUTS_DIR, exist_ok=True)

device = "cuda" if torch.cuda.is_available() else "cpu"
print(f"Device : {device}")

latest = best_checkpoint(cfg.CHECKPOINTS_D)
print(f"Meilleur checkpoint validation : {os.path.basename(latest)}")

pl_module = NLLBFineTuner.load_from_checkpoint(latest, cfg=cfg)
pl_module.eval().to(device)

tokenizer   = pl_module.tokenizer
model       = pl_module.model
tgt_lang_id = tokenizer.convert_tokens_to_ids(cfg.TGT_LANG)


def translate(text: str) -> str:
    if not text:
        return ""
    inputs = tokenizer(
        text, return_tensors="pt", padding=True,
        truncation=True, max_length=cfg.MAX_LENGTH,
    ).to(device)
    with torch.no_grad():
        generated = model.generate(
            **inputs,
            forced_bos_token_id=tgt_lang_id,
            max_new_tokens=cfg.MAX_LENGTH,
            num_beams=5,
        )
    return tokenizer.batch_decode(generated, skip_special_tokens=True)[0]


with open(os.path.join(cfg.DATA_DIR, "test.json"), encoding="utf-8") as f:
    samples = json.load(f)[:20]

doc = Document()
doc.add_heading("Resultats — Config D : NLLB-200 + LoRA + Back-Translation", 0)
doc.add_paragraph(
    "Identique a la Config C (LoRA), mais les donnees d'entrainement ont ete "
    "augmentées par rétro-traduction sérère → français. "
    "Les scores doivent provenir de evaluate_test.py; aucun score n'est codé en dur."
)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "Phrase source (francais)"
hdr[1].text = "Traduction predite (serere)"
hdr[2].text = "Reference (serere)"

print(f"\nTraduction de {len(samples)} exemples...\n")
for item in samples:
    src, ref = item["francais"], item["serere"]
    pred = translate(src)
    print(f"FR  : {src}\nSRR : {pred}\nREF : {ref}\n" + "-" * 50)
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text = src, pred, ref

output_path = os.path.join(cfg.OUTPUTS_DIR, "resultats_config_D.docx")
doc.save(output_path)
print(f"\nDocument Word cree : {output_path}")
