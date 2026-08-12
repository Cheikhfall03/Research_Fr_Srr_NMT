"""Génère un dossier de correction pour un second linguiste/annotateur.

Répond à la critique des trois reviewers sur l'évaluation humaine (un seul
annotateur, échantillon de 5-10 sorties par configuration): ce script produit
un lot élargi (par défaut 30 phrases par configuration, tirées aléatoirement
et de façon reproductible du test set, avec en option les 50 phrases UDHR de
calibration hors-domaine) et l'exporte à la fois en .docx prêt à imprimer/
annoter et en .csv prêt à dépouiller, afin qu'un deuxième correcteur puisse
noter les traductions indépendamment du premier.

Prérequis: outputs/predictions_<CONFIG>[.jsonl] doit exister (lancer
evaluate_test.py d'abord). Si des prédictions sont absentes pour une
configuration, les lignes correspondantes sont générées avec la traduction
laissée en blanc et un avertissement est affiché — le document reste
utilisable pour calibrer le format avec le correcteur en attendant.

Usage:
    python build_correctors_package.py --configs B C D E --n-samples 30
    python build_correctors_package.py --configs C D --include-udhr
"""
from __future__ import annotations

import argparse
import csv
import json
import random
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from config import ProbeConfig

OUTPUT_DIR = Path("annotations/round2_correctors")
CRITERIA = [
    ("nominal_class_1_5", "Accord nominal (1-5)"),
    ("consonant_mutation_1_5", "Mutation consonantique (1-5)"),
    ("implosives_1_5", "Implosives (1-5)"),
    ("apocope_1_5", "Apocope (1-5)"),
    ("syntax_1_5", "Syntaxe (1-5)"),
    ("semantic_fidelity_1_5", "Fidélité sémantique (1-5)"),
    ("overall_1_5", "Note globale (1-5)"),
]


def load_predictions(config: str, outputs_dir: Path, suffix: str) -> dict[int, str] | None:
    path = outputs_dir / f"predictions_{config}{suffix}.jsonl"
    if not path.exists():
        return None
    rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    return {row["id"]: row["prediction"] for row in rows}


def build_rows(config: str, sample_ids: list[int], test_rows: list[dict], predictions: dict[int, str] | None):
    rows = []
    for example_id in sample_ids:
        source_row = test_rows[example_id]
        prediction = predictions.get(example_id, "") if predictions else ""
        rows.append({
            "example_id": example_id,
            "configuration": config,
            "francais": source_row["francais"],
            "serere_reference": source_row["serere"],
            "prediction": prediction,
        })
    return rows


def write_csv(all_rows: list[dict], path: Path) -> None:
    fields = ["example_id", "configuration", "francais", "serere_reference", "prediction",
              "annotator_id", *[c for c, _ in CRITERIA], "target_language", "comment"]
    with open(path, "w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        for row in all_rows:
            writer.writerow({**{k: "" for k in fields}, **row})


def write_docx(all_rows: list[dict], path: Path, n_samples: int, configs: list[str], missing: list[str]) -> None:
    doc = Document()

    title = doc.add_heading("Fiche de correction — second annotateur", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("Traduction Automatique Neurale Français → Sérère")
    doc.add_paragraph("Document destiné à un second linguiste, en complément de la première évaluation "
                       "experte, afin d'élargir l'échantillon noté et de permettre le calcul d'un accord "
                       "inter-annotateurs.")
    meta = doc.add_paragraph()
    meta.add_run("Correcteur : ______________________     ").bold = False
    meta.add_run("Date : ______________________")

    doc.add_heading("Instructions", level=2)
    instructions = doc.add_paragraph()
    instructions.add_run(
        f"Pour chaque phrase ci-dessous, comparez la traduction du système à la traduction de "
        f"référence, puis notez chaque critère de 1 (très mauvais) à 5 (excellent) dans la colonne "
        f"correspondante. Échantillon : {n_samples} phrases par configuration "
        f"({', '.join(configs)}), tirées aléatoirement du test set (seed fixe pour reproductibilité)."
    )
    if missing:
        warn = doc.add_paragraph()
        run = warn.add_run(
            "ATTENTION : traductions manquantes pour les configurations suivantes (à régénérer avec "
            f"evaluate_test.py avant impression) : {', '.join(missing)}."
        )
        run.bold = True
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    doc.add_heading("Critères", level=2)
    for _, label in CRITERIA[:-1]:
        doc.add_paragraph(label, style="List Bullet")
    doc.add_paragraph(
        "Type de langue cible produite (target_language) : indiquez si la sortie est bien du sérère, "
        "ou si elle dérive vers un autre idiome (ex. wolof) — pertinent pour la question de la "
        "contamination inter-langues discutée dans l'article."
    )

    for config in configs:
        doc.add_heading(f"Configuration {config}", level=2)
        table = doc.add_table(rows=1, cols=5 + len(CRITERIA) + 1)
        table.style = "Light Grid Accent 1"
        header_cells = table.rows[0].cells
        headers = ["ID", "Français", "Référence sérère", "Traduction système", "Langue cible obtenue",
                   *[label for _, label in CRITERIA], "Commentaire"]
        for cell, text in zip(header_cells, headers):
            cell.text = text
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(9)

        for row in [r for r in all_rows if r["configuration"] == config]:
            cells = table.add_row().cells
            values = [str(row["example_id"]), row["francais"], row["serere_reference"], row["prediction"],
                      "", *["" for _ in CRITERIA], ""]
            for cell, text in zip(cells, values):
                cell.text = text
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(9)
        doc.add_paragraph()

    doc.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--configs", nargs="+", choices=list("ABCDE"), default=list("BCDE"))
    parser.add_argument("--n-samples", type=int, default=30)
    parser.add_argument("--seed", type=int, default=42, help="Seed du tirage aléatoire des phrases échantillonnées.")
    parser.add_argument("--outputs-dir", default="outputs")
    parser.add_argument("--suffix", default="", help="Suffixe des fichiers de prédictions, ex. _seed42.")
    parser.add_argument("--include-udhr", action="store_true",
                         help="Ajoute les 50 phrases UDHR (calibration hors-domaine) en plus de l'échantillon test.")
    args = parser.parse_args()

    base = ProbeConfig()
    test_rows = json.loads((Path(base.DATA_DIR) / "test.json").read_text(encoding="utf-8"))
    rng = random.Random(args.seed)
    sample_ids = sorted(rng.sample(range(len(test_rows)), min(args.n_samples, len(test_rows))))

    outputs_dir = Path(args.outputs_dir)
    all_rows, missing = [], []
    for config in args.configs:
        predictions = load_predictions(config, outputs_dir, args.suffix)
        if predictions is None:
            missing.append(config)
        all_rows.extend(build_rows(config, sample_ids, test_rows, predictions))

    if args.include_udhr:
        udhr_rows = json.loads((Path(base.DATA_DIR) / "udhr_fr_srr_probe_test.json").read_text(encoding="utf-8"))
        calibration_dir = Path("results/calibration_udhr")
        for config in args.configs:
            path = calibration_dir / f"predictions_{config}{args.suffix}.jsonl"
            predictions = None
            if path.exists():
                rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
                predictions = {row["id"]: row["prediction"] for row in rows}
            else:
                missing.append(f"{config} (UDHR)")
            all_rows.extend(build_rows(config, list(range(len(udhr_rows))), udhr_rows, predictions))

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    write_csv(all_rows, OUTPUT_DIR / f"scoring_template{args.suffix}.csv")
    write_docx(all_rows, OUTPUT_DIR / f"correctors_package{args.suffix}.docx",
               args.n_samples, args.configs, sorted(set(missing)))

    print(f"{len(all_rows)} lignes générées pour {len(args.configs)} configuration(s).")
    if missing:
        print(f"ATTENTION — prédictions manquantes pour : {sorted(set(missing))}. "
              "Relancez evaluate_test.py / evaluate_udhr_calibration.py puis regénérez ce package.")
    print(f"Package -> {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
