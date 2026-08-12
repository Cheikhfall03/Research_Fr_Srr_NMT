"""
Inference — Baseline Transformer from scratch.
Traduit 20 exemples du split test → outputs/resultats_baseline_scratch.docx
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import json
import torch
from docx import Document
from config import ScratchConfig
from reproducibility import best_checkpoint
from baselines.models.model_scratch import ScratchTransformer

cfg = ScratchConfig()
os.makedirs(cfg.OUTPUTS_DIR, exist_ok=True)

device = "cuda" if torch.cuda.is_available() else "cpu"
print(f"Device : {device}")

CKPT_DIR = os.path.join(cfg.CHECKPOINTS_DIR, "baseline_scratch")
best = best_checkpoint(CKPT_DIR)
print(f"Checkpoint : {os.path.basename(best)}")

pl_module = ScratchTransformer.load_from_checkpoint(best, cfg=cfg)
pl_module.eval().to(device)
tokenizer = pl_module.tokenizer
model     = pl_module.model


def translate(text: str) -> str:
    if not text:
        return ""
    ids     = tokenizer.encode(text, cfg.MAX_LENGTH)
    padded, mask = tokenizer.pad([ids], cfg.MAX_LENGTH)
    src  = torch.tensor(padded, dtype=torch.long).to(device)
    amsk = torch.tensor(mask,   dtype=torch.long).to(device)
    with torch.no_grad():
        gen = model.generate(src, amsk, max_new_tokens=cfg.MAX_LENGTH, num_beams=1)
    return tokenizer.decode(gen[0].tolist())


with open(os.path.join(cfg.DATA_DIR, "test.json"), encoding="utf-8") as f:
    samples = json.load(f)[:20]

doc = Document()
doc.add_heading("Resultats — Baseline : Transformer from scratch", 0)
doc.add_paragraph(
    "Transformer standard (PyTorch nn.Transformer, d_model=128, 3 couches enc/dec) "
    "entraîné from scratch sur 23 113 paires FR→SRR. "
    "Tokenizer SentencePiece BPE vocab=8000 construit depuis les données."
)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "Phrase source (francais)"
hdr[1].text = "Traduction predite (serere)"
hdr[2].text = "Reference (serere)"

print(f"\nTraduction de {len(samples)} exemples...\n")
for item in samples:
    src, ref = item["francais"], item["serere"]
    pred = translate(src)
    print(f"FR  : {src}\nSRR : {pred}\nREF : {ref}\n" + "-" * 50)
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text = src, pred, ref

output_path = os.path.join(cfg.OUTPUTS_DIR, "resultats_baseline_scratch.docx")
doc.save(output_path)
print(f"\nDocument Word créé : {output_path}")
