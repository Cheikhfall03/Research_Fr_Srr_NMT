"""
Inference — Baseline opus-mt-fr fine-tuné.
Traduit 20 exemples du split test → outputs/resultats_baseline_opusmt.docx
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import json
import torch
from docx import Document
from config import OpusMTConfig
from reproducibility import best_checkpoint
from baselines.models.model_opusmt import OpusMTFineTuner

cfg = OpusMTConfig()
os.makedirs(cfg.OUTPUTS_DIR, exist_ok=True)

device = "cuda" if torch.cuda.is_available() else "cpu"
print(f"Device : {device}")

CKPT_DIR = os.path.join(cfg.CHECKPOINTS_DIR, "baseline_opusmt")
best = best_checkpoint(CKPT_DIR)
print(f"Checkpoint : {os.path.basename(best)}")

pl_module = OpusMTFineTuner.load_from_checkpoint(best, cfg=cfg)
pl_module.eval().to(device)
source_tokenizer = pl_module.source_tokenizer
target_tokenizer = pl_module.target_tokenizer
model     = pl_module.model


def translate(text: str) -> str:
    if not text:
        return ""
    inputs = source_tokenizer(text, return_tensors="pt", truncation=True, max_length=cfg.MAX_LENGTH).to(device)
    with torch.no_grad():
        generated = model.generate(**inputs, max_new_tokens=cfg.MAX_LENGTH, num_beams=5)
    return target_tokenizer.batch_decode(generated)[0]


with open(os.path.join(cfg.DATA_DIR, "test.json"), encoding="utf-8") as f:
    samples = json.load(f)[:20]

doc = Document()
doc.add_heading("Resultats — Baseline : opus-mt-fr fine-tuné", 0)
doc.add_paragraph(
    "Modèle de traduction pré-entraîné sur le français (Helsinki-NLP/opus-mt-fr-en), "
    "adapté au sérère par fine-tuning complet sur 23 113 paires FR→SRR."
)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "Phrase source (francais)"
hdr[1].text = "Traduction predite (serere)"
hdr[2].text = "Reference (serere)"

print(f"\nTraduction de {len(samples)} exemples...\n")
for item in samples:
    src, ref = item["francais"], item["serere"]
    pred = translate(src)
    print(f"FR  : {src}\nSRR : {pred}\nREF : {ref}\n" + "-" * 50)
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text = src, pred, ref

output_path = os.path.join(cfg.OUTPUTS_DIR, "resultats_baseline_opusmt.docx")
doc.save(output_path)
print(f"\nDocument Word créé : {output_path}")
